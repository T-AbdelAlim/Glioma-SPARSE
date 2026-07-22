r"""
End-to-end Glioma-SPARSE inference: Stage A grade + Stage B molecular subtype.

Input is a single WSI or a folder of WSIs (.ndpi, .svs, .mrxs). For each slide:

  1. generate a 2048x2048 thumbnail and its JSON mapping sidecar,
  2. Stage A grade prediction, patch-injection risk map, and high-risk tile
     extraction from the original WSI via the sidecar (p95 by default,
     configurable to p90/p97/etc.),
  3. Stage B molecular subtype prediction on the extracted patches, soft-voted
     to slide level,
  4. an integrated WHO 2021 diagnosis report (.xlsx) with per-stage predictions,
     probabilities, and confidence, including a confidence for the integrated
     call, coloured by uncertainty,
  5. optional Stage B occlusion risk map: each tile of a Stage-B patch is
     occluded and the drop in the predicted-class logit is recorded; the
     strongest tiles are re-extracted from the WSI at 2048x2048 (a third,
     cellular-level zoom) and saved as JPGs so the IDH-indicative regions can
     be inspected,
  6. a testset mode that runs each fold's held-out test slides through that
     fold's models, evaluating controls fairly and reporting subtype accuracy,
     a wildtype-vs-mutant accuracy, and a separate 1p/19q codeletion table,
     aggregated across folds.

All outputs are written under:
    Glioma-SPARSE/results/<run-name>/<slide-abbrev>/{riskmap,patches}/

Run modes:
  * Edit the CONFIG block below and press Run (PyCharm), or
  * pass command-line flags (these override CONFIG), e.g.

    python -m scripts.inference.inference_end_to_end \
        --input path/to/slide.ndpi \
        --stage-a-ckpt training_output/<fold>/best_auc.pth \
        --stage-b-ckpt training_output_stageB/<fold>/best_auc.pth \
        --run-name e2e_single [--percentile 95] [--stage-b-riskmap]

    python -m scripts.inference.inference_end_to_end --testset \
        --splits-dir splits \
        --stage-a-glob "training_output/*_split_0*/best_auc.pth" \
        --stage-b-glob "training_output_stageB/*_fold*/best_auc.pth" \
        --sidecar-root data/ebrains_thumbnails \
        --wsi-root path/to/WHO2021_data --run-name e2e_testset

e.g. split 2 testset: python -m scripts.inference.inference_end_to_end --testset --splits-dir splits --stage-a-glob "training_output/20260703_0554_resnet18_cw_split_02/best_auc.pth" --stage-b-glob "training_output_stageB/20260705_1203_resnet18_stageB_fold2_cw/best_auc.pth" --sidecar-root "data/ebrains_thumbnails" --wsi-root "C:\Users\TAbde\Documents\EMC_postdoc\Virtual_Biopsy\data\WSI_ebrains\WHO2021_data" --run-name e2e_testset_split_2

"""

import argparse
import csv
import glob
import sys
from collections import defaultdict
from pathlib import Path

import numpy as np
from PIL import Image
import torch

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

from glioma_sparse.preprocessing.create_wsi_thumbnail import create_wsi_thumbnail
from glioma_sparse.data_utils.transforms import build_eval_transform
from glioma_sparse.models.factory import build_model
from glioma_sparse.interpret.patch_injection import (
    compute_risk_map, tile_image, reconstruct_image,
)
from glioma_sparse.interpret.wsi_mapping import (
    load_mapping, thumbnail_bbox_to_wsi, grid_cell_bbox,
)
import openslide


# ============================================================
# CONFIG  (edit these, then press Run; CLI flags override them)
# ============================================================

REPO_ROOT = Path(__file__).resolve().parents[2]

CONFIG = {
    # what to run: "single", "folder", or "testset"
    "run_mode": "folder",
    "run_name": "external_TCGA_val/ext_TCGA_astro_G2",

    # single / folder mode
    "input": r"E:\Thinkpad_Backup\Data\WSI_datasets\TCGA_data_download\manifest_grades\TCGA_ext_val\low_grade\astro_IDHmt_G2",
    "stage_a_ckpt": r"C:\Users\TAbde\PycharmProjects\Glioma-SPARSE\ResNet18_output\training_output\20260703_0554_resnet18_cw_split_02/best_auc.pth",
    "stage_b_ckpt": r"C:\Users\TAbde\PycharmProjects\Glioma-SPARSE\ResNet18_output\training_output_stageB\20260705_1203_resnet18_stageB_fold2_cw/best_auc.pth",

    # testset mode (per-fold, leakage-safe)
    "splits_dir": "splits",
    "stage_a_glob": "ResNet18_output/training_output/20260703_0554_resnet18_cw_split_02/best_auc.pth",
    "stage_b_glob": "ResNet18_output/training_output_stageB/20260705_1203_resnet18_stageB_fold2_cw/best_auc.pth",
    "sidecar_root": r"data\ebrains_thumbnails",

    # shared
    "wsi_root": r"C:\Users\TAbde\Documents\EMC_postdoc\Virtual_Biopsy\data\WSI_ebrains\WHO2021_data",
    "control_image": r"C:\Users\TAbde\PycharmProjects\Glioma-SPARSE\data\included\control\86242943-7775-11eb-827d-001a7dda7111.jpg" ,
    "model": "resnet18",

    # region selection
    "percentile": 95.0,        # p95 default; set 90 or 97 to change
    "cap": 4,
    "min_tissue": 0.25,         # default was 0.5 during training

    # thumbnailing / patches
    "target_mpp": 4.0,
    "thumb_size": 2048,
    "output_size": 2048,

    # Stage B occlusion risk map + cellular zoom
    "stage_b_riskmap": True,
    "occlusion_topk": 1,       # strongest occlusion tiles re-zoomed per patch

    "device": "cuda" if torch.cuda.is_available() else "cpu",
}

GRADE_CLASSES = ["control", "low_grade", "high_grade"]
SUBTYPE_CLASSES = ["IDH_mt", "IDH_mt_1p19q", "IDH_wt"]
MUTANT_SET = {"IDH_mt", "IDH_mt_1p19q"}
GRID = (8, 8)
SEED = 42
WSI_EXTS = {".ndpi", ".svs", ".mrxs", ".tif", ".tiff"}
RESULTS_ROOT = REPO_ROOT / "results"

# Figure export: 300 dpi JPG, sized so the saved image lands around 2048 px.
FIG_DPI = 300
FIG_INCHES = 2048 / FIG_DPI     # ~6.83 in -> ~2048 px at 300 dpi
JPG_DPI = (300, 300)
JPG_QUALITY = 90

INTEGRATED = {
    ("low", "IDH_mt"): "Astrocytoma, IDH-mutant, grade 2",
    ("high", "IDH_mt"): "Astrocytoma, IDH-mutant, grade 3-4",
    ("low", "IDH_mt_1p19q"): "Oligodendroglioma, IDH-mutant 1p/19q-codeleted, grade 2",
    ("high", "IDH_mt_1p19q"): "Oligodendroglioma, IDH-mutant 1p/19q-codeleted, grade 3",
    ("low", "IDH_wt"): "IDH-wildtype glioma (low-grade morphology)",
    ("high", "IDH_wt"): "Glioblastoma, IDH-wildtype, grade 4",
}


# ============================================================
# MODELS + CONFIDENCE
# ============================================================

def load_model(ckpt, model_name, n_classes, device):
    model = build_model(model_name, num_classes=n_classes)
    state = torch.load(str(ckpt), map_location=device, weights_only=True)
    model.load_state_dict(state)
    model.to(device).eval()
    return model


def forward_probs(model, transform, img, device):
    x = transform(img).unsqueeze(0).to(device)
    with torch.no_grad():
        logits = model(x)
        probs = torch.softmax(logits, dim=1).cpu().numpy()[0]
    return probs, logits.cpu().numpy()[0]


def confidence(probs):
    p = np.asarray(probs, dtype=float)
    order = np.argsort(p)[::-1]
    top = float(p[order[0]])
    second = float(p[order[1]]) if len(p) > 1 else 0.0
    K = len(p)
    ent = -np.sum(p * np.log(p + 1e-12))
    norm_ent = float(ent / np.log(K))
    return {"pred_idx": int(order[0]), "top_prob": round(top, 4),
            "margin": round(top - second, 4),
            "confidence": round(1.0 - norm_ent, 4)}


def grade_to_class(grade_name):
    return {"low_grade": "low", "high_grade": "high"}.get(grade_name, "control")


def mutant_status(subtype):
    return "mutant" if subtype in MUTANT_SET else "wildtype"


def abbrev(slide_id):
    """Short, human-readable folder name from a slide id (UUID-friendly)."""
    head = slide_id.split("-")[0]
    return head if len(head) >= 6 else slide_id[:12]


# ============================================================
# STAGE A: region selection + high-res extraction
# ============================================================

def select_tiles(risk_map, thumbnail, grid, percentile, cap, min_tissue):
    rows, cols = grid
    thresh = float(np.percentile(risk_map.flatten(), percentile))
    kept = []
    for r in range(rows):
        for c in range(cols):
            if risk_map[r, c] >= thresh:
                tf = _tissue_fraction(thumbnail, r, c, grid)
                if tf >= min_tissue:
                    kept.append({"row": r, "col": c, "risk": float(risk_map[r, c]),
                                 "tissue_fraction": tf})
    kept.sort(key=lambda d: -d["risk"])
    return kept[:cap], thresh


def _tissue_fraction(thumbnail, row, col, grid):
    rows, cols = grid
    w, h = thumbnail.size
    tw, th = w // cols, h // rows
    tile = thumbnail.crop((col * tw, row * th, col * tw + tw, row * th + th))
    arr = np.array(tile).astype(np.float32) / 255.0
    maxc, minc = arr.max(axis=2), arr.min(axis=2)
    sat = (maxc - minc) / (maxc + 1e-8)
    mask = (maxc < 0.9) & (sat > 0.05)
    return float(mask.sum()) / float(mask.size)


def extract_highres(slide, mapping, row, col, grid, out_size):
    bbox_thumb = grid_cell_bbox(row, col, grid[0], grid[1], mapping.thumbnail_size)
    wsi_bbox = thumbnail_bbox_to_wsi(mapping, bbox_thumb)
    if wsi_bbox is None:
        return None, None
    wx, wy, ww, wh = wsi_bbox
    patch = slide.read_region((wx, wy), 0, (ww, wh)).convert("RGB")
    if patch.size != (out_size, out_size):
        patch = patch.resize((out_size, out_size), Image.BICUBIC)
    return patch, wsi_bbox


def subtile_wsi_bbox(patch_wsi_bbox, sr, sc, grid=8):
    """WSI level-0 bbox of occlusion sub-tile (sr,sc) within a Stage-B patch."""
    wx, wy, ww, wh = patch_wsi_bbox
    sub_ww, sub_wh = ww // grid, wh // grid
    return (wx + sc * sub_ww, wy + sr * sub_wh, sub_ww, sub_wh)


def extract_subtile(slide, patch_wsi_bbox, sr, sc, out_size, grid=8):
    sx, sy, sw, sh = subtile_wsi_bbox(patch_wsi_bbox, sr, sc, grid)
    if sw <= 0 or sh <= 0:
        return None
    tile = slide.read_region((sx, sy), 0, (sw, sh)).convert("RGB")
    return tile.resize((out_size, out_size), Image.BICUBIC)


# ============================================================
# STAGE B: occlusion importance
# ============================================================

def occlusion_map(model, transform, patch, pred_idx, device, grid=GRID):
    """Occlude each tile; importance = base_logit - occluded_logit."""
    rows, cols = grid
    _, base_logits = forward_probs(model, transform, patch, device)
    base = float(base_logits[pred_idx])
    tiles = tile_image(patch, grid)
    mean_col = tuple(int(v) for v in np.array(patch).reshape(-1, 3).mean(0))
    imp = np.zeros((rows, cols), dtype=np.float32)
    for r in range(rows):
        for c in range(cols):
            occ = list(tiles)
            occ[r * cols + c] = Image.new("RGB", occ[r * cols + c].size, mean_col)
            occ_img = reconstruct_image(occ, grid, patch.size)
            _, zl = forward_probs(model, transform, occ_img, device)
            imp[r, c] = base - float(zl[pred_idx])
    return imp


def top_occlusion_cells(imp, k):
    flat = [(imp[r, c], r, c) for r in range(imp.shape[0]) for c in range(imp.shape[1])]
    flat.sort(key=lambda t: -t[0])
    return [(r, c, float(v)) for v, r, c in flat[:k] if v > 0]


# ============================================================
# VISUALISATION
# ============================================================

def _overlay(ax, image, heat, grid, boxes=None, box_color="#e11d1d"):
    ax.imshow(image)
    rows, cols = grid
    w, h = image.size
    tw, th = w / cols, h / rows
    hn = (heat - heat.min()) / (np.ptp(heat) + 1e-8)
    ax.imshow(np.kron(hn, np.ones((int(th), int(tw)))),
              cmap="RdYlBu_r", alpha=0.45, extent=[0, w, h, 0])
    for r in range(rows + 1):
        ax.axhline(r * th, color="white", lw=0.4, alpha=0.5)
    for c in range(cols + 1):
        ax.axvline(c * tw, color="white", lw=0.4, alpha=0.5)
    for (r, c) in (boxes or []):
        ax.add_patch(plt.Rectangle((c * tw, r * th), tw, th, fill=False,
                                   edgecolor=box_color, lw=2.2))
    ax.set_xticks([]); ax.set_yticks([])


def save_stage_a_riskmap(thumbnail, risk_map, selected, path, title):
    fig, ax = plt.subplots(figsize=(FIG_INCHES, FIG_INCHES))
    boxes = [(s["row"], s["col"]) for s in selected]
    _overlay(ax, thumbnail, risk_map, GRID, boxes)
    ax.set_title(title, fontsize=10)
    fig.tight_layout()
    fig.savefig(path, dpi=FIG_DPI, format="jpg", bbox_inches="tight",
                pil_kwargs={"quality": JPG_QUALITY})
    plt.close(fig)


def save_stage_b_occlusion(patch, imp, top_cells, path, title):
    fig, ax = plt.subplots(figsize=(FIG_INCHES, FIG_INCHES))
    boxes = [(r, c) for r, c, _ in top_cells]
    _overlay(ax, patch, imp, GRID, boxes)
    ax.set_title(title, fontsize=10)
    fig.tight_layout()
    fig.savefig(path, dpi=FIG_DPI, format="jpg", bbox_inches="tight",
                pil_kwargs={"quality": JPG_QUALITY})
    plt.close(fig)


# ============================================================
# PER-SLIDE PIPELINE
# ============================================================

def run_slide(thumb_source, model_a, model_b, transform, control_img, args,
              slide_dir, wsi_path=None, force_skip_stage_b=False,
              generate_thumb=False):
    """Full Stage A (+ Stage B) pass for one slide. Writes risk maps and patches
    under slide_dir. Returns a result dict."""
    slide_id = Path(thumb_source).stem
    riskmap_dir = slide_dir / "riskmap"
    patches_dir = slide_dir / "patches"
    riskmap_dir.mkdir(parents=True, exist_ok=True)

    # 1. thumbnail (+ sidecar). Generated fresh for deployment, reused otherwise.
    thumb_path = Path(thumb_source)
    if generate_thumb:
        thumb_path = slide_dir / f"{slide_id}.jpg"
        create_wsi_thumbnail(wsi_path, output_path=str(thumb_path),
                             target_mpp=args.target_mpp, output_size=args.thumb_size)
    thumbnail = Image.open(thumb_path).convert("RGB")

    # 2. Stage A grade + risk map
    probsA, _ = forward_probs(model_a, transform, thumbnail, args.device)
    cA = confidence(probsA)
    grade_name = GRADE_CLASSES[cA["pred_idx"]]
    grade_class = grade_to_class(grade_name)

    risk_map, _ = compute_risk_map(
        target_img=thumbnail, control_img=control_img, model=model_a,
        transform=transform, grid=GRID, target_class=cA["pred_idx"],
        device=args.device, control_shuffle_seed=SEED, score="logit")
    selected, thresh = select_tiles(risk_map, thumbnail, GRID,
                                    args.percentile, args.cap, args.min_tissue)
    save_stage_a_riskmap(thumbnail, risk_map, selected,
                         riskmap_dir / "stageA_riskmap.jpg",
                         f"{slide_id}  Stage A: {grade_name} ({cA['confidence']:.2f})")

    result = {
        "slide_id": slide_id, "grade_pred": grade_name, "grade_class": grade_class,
        "prob_control": round(float(probsA[0]), 4),
        "prob_low_grade": round(float(probsA[1]), 4),
        "prob_high_grade": round(float(probsA[2]), 4),
        "stageA_confidence": cA["confidence"], "stageA_margin": cA["margin"],
        "p_threshold": round(thresh, 6), "n_regions": len(selected),
    }

    # control (predicted or, in testset, known) stops after Stage A
    if grade_name == "control" or force_skip_stage_b or not selected:
        diag = ("Control tissue" if grade_name == "control"
                else "No molecular subtyping" if force_skip_stage_b
                else "No region extracted")
        result.update({"subtype_pred": "NA", "integrated_diagnosis": diag,
                       "stageB_confidence": "", "integrated_confidence": ""})
        return result, risk_map

    # 3. Stage B on extracted high-res patches
    mapping = load_mapping(thumb_path)
    if not mapping.has_wsi_mapping():
        raise RuntimeError(f"no WSI mapping for {slide_id}")
    wsi_real = resolve_wsi(mapping.wsi_path, args, wsi_path)
    slide = openslide.OpenSlide(wsi_real)
    patches_dir.mkdir(parents=True, exist_ok=True)
    patch_probs = []
    try:
        for rank, sel in enumerate(selected, 1):
            patch, patch_bbox = extract_highres(slide, mapping, sel["row"],
                                                sel["col"], GRID, args.output_size)
            if patch is None:
                continue
            tag = f"{slide_id}_rank{rank:02d}_r{sel['row']}c{sel['col']}"
            patch.save(patches_dir / f"{tag}.jpg", quality=JPG_QUALITY, dpi=JPG_DPI)
            pB, _ = forward_probs(model_b, transform, patch, args.device)
            patch_probs.append(pB)

            # 5. optional occlusion risk map + cellular zoom of top tiles
            if args.stage_b_riskmap:
                imp = occlusion_map(model_b, transform, patch,
                                    int(np.argmax(pB)), args.device)
                tops = top_occlusion_cells(imp, args.occlusion_topk)
                save_stage_b_occlusion(
                    patch, imp, tops, riskmap_dir / f"stageB_{tag}_occlusion.jpg",
                    f"{tag}  Stage B occlusion ({SUBTYPE_CLASSES[int(np.argmax(pB))]})")
                for ti, (sr, sc, val) in enumerate(tops, 1):
                    zoom = extract_subtile(slide, patch_bbox, sr, sc, args.output_size)
                    if zoom is not None:
                        zoom.save(riskmap_dir /
                                  f"stageB_{tag}_top{ti}_r{sr}c{sc}_zoom.jpg",
                                  quality=JPG_QUALITY, dpi=JPG_DPI)
    finally:
        slide.close()

    if not patch_probs:
        result.update({"subtype_pred": "NA",
                       "integrated_diagnosis": "No region extracted",
                       "stageB_confidence": "", "integrated_confidence": ""})
        return result, risk_map

    # 4. soft-vote to slide-level subtype
    mean_prob = np.mean(patch_probs, axis=0)
    cB = confidence(mean_prob)
    subtype = SUBTYPE_CLASSES[cB["pred_idx"]]
    integrated = INTEGRATED.get((grade_class, subtype), "Unresolved")
    result.update({
        "subtype_pred": subtype,
        "prob_IDH_mt": round(float(mean_prob[0]), 4),
        "prob_IDH_mt_1p19q": round(float(mean_prob[1]), 4),
        "prob_IDH_wt": round(float(mean_prob[2]), 4),
        "stageB_confidence": cB["confidence"], "stageB_margin": cB["margin"],
        "integrated_diagnosis": integrated,
        "integrated_confidence": round(cA["confidence"] * cB["confidence"], 4),
        "n_patches_used": len(patch_probs),
    })
    return result, risk_map


def resolve_wsi(stored, args, explicit=None):
    if explicit and Path(explicit).exists():
        return explicit
    if stored and Path(stored).exists():
        return stored
    if args.wsi_root:
        hit = _wsi_index(args.wsi_root).get(Path(str(stored).replace("\\", "/")).name)
        if hit:
            return hit
    raise FileNotFoundError(f"WSI not found: {stored}")


_WSI_IDX = {}
def _wsi_index(root):
    if root not in _WSI_IDX:
        idx = {}
        for p in Path(root).rglob("*"):
            if p.suffix.lower() in WSI_EXTS:
                idx.setdefault(p.name, str(p))
        _WSI_IDX[root] = idx
    return _WSI_IDX[root]


# ============================================================
# REPORTING (xlsx, coloured by uncertainty)
# ============================================================

REPORT_COLS = ["slide_id", "fold", "grade_pred", "grade_class",
               "prob_control", "prob_low_grade", "prob_high_grade",
               "stageA_confidence", "subtype_pred",
               "prob_IDH_mt", "prob_IDH_mt_1p19q", "prob_IDH_wt",
               "stageB_confidence", "integrated_diagnosis", "integrated_confidence",
               "true_grade_class", "true_mutation",
               "grade_correct", "subtype_correct", "mutant_status_correct",
               "end_to_end_correct"]


def _conf_fill(value):
    from openpyxl.styles import PatternFill
    try:
        v = float(value)
    except (TypeError, ValueError):
        return None
    if v >= 0.66:
        c = "C6EFCE"       # green
    elif v >= 0.33:
        c = "FFEB9C"       # amber
    else:
        c = "FFC7CE"       # red
    return PatternFill("solid", fgColor=c)


def write_xlsx_report(rows, path, extra_sheets=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = "per_slide"
    cols = [c for c in REPORT_COLS if any(c in r for r in rows)]
    ws.append(cols)
    head_fill = PatternFill("solid", fgColor="305496")
    for j, _ in enumerate(cols, 1):
        cell = ws.cell(row=1, column=j)
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center")
    conf_cols = {"stageA_confidence", "stageB_confidence", "integrated_confidence"}
    for i, r in enumerate(rows, 2):
        for j, col in enumerate(cols, 1):
            val = r.get(col, "")
            cell = ws.cell(row=i, column=j, value=val)
            if col in conf_cols:
                fill = _conf_fill(val)
                if fill:
                    cell.fill = fill
    for j, col in enumerate(cols, 1):
        width = max(len(col), *(len(str(r.get(col, ""))) for r in rows)) + 2
        ws.column_dimensions[ws.cell(row=1, column=j).column_letter].width = min(width, 40)
    ws.freeze_panes = "A2"

    for name, (header, table) in (extra_sheets or {}).items():
        s = wb.create_sheet(name)
        s.append(header)
        for j in range(1, len(header) + 1):
            c = s.cell(row=1, column=j)
            c.font = Font(bold=True, color="FFFFFF"); c.fill = head_fill
        for row in table:
            s.append(row)
        for j, h in enumerate(header, 1):
            width = max(len(str(h)), *(len(str(row[j-1])) for row in table)) + 2 if table else len(str(h)) + 2
            s.column_dimensions[s.cell(row=1, column=j).column_letter].width = min(width, 44)
    wb.save(path)


# ============================================================
# DRIVERS
# ============================================================

def list_wsis(input_path):
    p = Path(input_path)
    return [p] if p.is_file() else sorted(
        q for q in p.rglob("*") if q.suffix.lower() in WSI_EXTS)


def run_single_or_folder(args):
    transform = build_eval_transform()
    control_img = Image.open(_rel(args.control_image)).convert("RGB")
    model_a = load_model(_rel(args.stage_a_ckpt), args.model,
                         len(GRADE_CLASSES), args.device)
    model_b = load_model(_rel(args.stage_b_ckpt), args.model,
                         len(SUBTYPE_CLASSES), args.device)
    run_dir = RESULTS_ROOT / args.run_name
    run_dir.mkdir(parents=True, exist_ok=True)

    rows = []
    for wsi in list_wsis(_rel(args.input)):
        slide_id = wsi.stem
        slide_dir = run_dir / abbrev(slide_id)
        try:
            res, _ = run_slide(wsi, model_a, model_b, transform, control_img, args,
                               slide_dir, wsi_path=str(wsi), generate_thumb=True)
            rows.append(res)
            print(f"  {slide_id}: {res['integrated_diagnosis']} "
                  f"(conf {res.get('integrated_confidence', '')})")
        except Exception as e:
            print(f"  FAILED {wsi.name}: {e}")
    write_xlsx_report(rows, run_dir / "integrated_diagnosis_report.xlsx")
    print(f"\nReport: {run_dir / 'integrated_diagnosis_report.xlsx'}")


def _rel(path):
    """Resolve a path against the repo root when it is relative, so the script
    works whether it is run as a module or as a file (any working directory)."""
    p = Path(path)
    return p if p.is_absolute() else (REPO_ROOT / p)


def run_testset(args):
    transform = build_eval_transform()
    control_img = Image.open(_rel(args.control_image)).convert("RGB")
    run_dir = RESULTS_ROOT / args.run_name
    run_dir.mkdir(parents=True, exist_ok=True)

    split_csvs = sorted(_rel(args.splits_dir).glob("split_0*.csv"))
    a_ckpts = sorted(glob.glob(str(_rel(args.stage_a_glob))))
    b_ckpts = sorted(glob.glob(str(_rel(args.stage_b_glob))))
    if not split_csvs:
        raise SystemExit(f"No split_0*.csv found in {_rel(args.splits_dir)}. "
                         f"Check --splits-dir or the working directory.")
    if not (a_ckpts and b_ckpts):
        raise SystemExit(f"No checkpoints matched.\n  stage A glob: "
                         f"{_rel(args.stage_a_glob)} -> {len(a_ckpts)} hits\n  "
                         f"stage B glob: {_rel(args.stage_b_glob)} -> {len(b_ckpts)} hits")

    def pick(ckpts, token):
        m = [c for c in ckpts if token in str(c)]
        return m[0] if m else None

    all_rows = []
    for fold_idx, sc in enumerate(split_csvs, 1):
        a = pick(a_ckpts, sc.stem); b = pick(b_ckpts, f"fold{fold_idx}")
        if not (a and b):
            print(f"[fold {fold_idx}] missing checkpoint, skip"); continue
        model_a = load_model(a, args.model, len(GRADE_CLASSES), args.device)
        model_b = load_model(b, args.model, len(SUBTYPE_CLASSES), args.device)
        test = [Path(r["path"]) for r in csv.DictReader(open(sc))
                if r["split"] == "test"]
        print(f"[fold {fold_idx}] {len(test)} test slides")
        fold_dir = run_dir / f"fold_{fold_idx}"

        for split_path in test:
            stem = split_path.stem
            true_grade_folder = split_path.parent.name       # control/low_grade/high_grade
            is_control = (true_grade_folder == "control")
            if is_control:
                true = {"grade_class": "control", "mutation": "NA"}
                thumb_source = split_path                     # grade thumbnail (no sidecar)
            else:
                sc_thumb = _sidecar_for(stem, _rel(args.sidecar_root))
                if sc_thumb is None:
                    print(f"    no sidecar for {stem}"); continue
                true = _true_labels(sc_thumb)
                thumb_source = sc_thumb
            try:
                res, _ = run_slide(thumb_source, model_a, model_b, transform,
                                   control_img, args, fold_dir / abbrev(stem),
                                   force_skip_stage_b=is_control)
                res["fold"] = fold_idx
                res.update(score_row(true, res))
                all_rows.append(res)
            except Exception as e:
                print(f"    FAILED {stem}: {e}")

    _write_testset_outputs(all_rows, run_dir)
    print(f"\nTestset report: {run_dir / 'e2e_testset_report.xlsx'}")


def score_row(true, res):
    """Correctness flags with fair control handling."""
    tgc, tmut = true["grade_class"], true["mutation"]
    pred_gc, pred_sub = res["grade_class"], res.get("subtype_pred", "NA")
    grade_correct = int(pred_gc == tgc)
    out = {"true_grade_class": tgc, "true_mutation": tmut,
           "grade_correct": grade_correct}
    if tgc == "control":
        out.update({"stage_b_evaluated": 0, "subtype_correct": "",
                    "mutant_status_correct": "",
                    "end_to_end_correct": grade_correct})
        return out
    out["stage_b_evaluated"] = 1
    out["subtype_correct"] = int(pred_sub == tmut)
    out["mutant_status_correct"] = int(mutant_status(pred_sub) == mutant_status(tmut))
    out["end_to_end_correct"] = int(grade_correct and pred_sub == tmut)
    return out


def _sidecar_for(stem, root):
    for jpg in Path(root).rglob(f"{stem}.jpg"):
        if jpg.with_suffix(".json").exists():
            return jpg
    return None


def _true_labels(sc_thumb):
    folder = Path(sc_thumb).parent.parent.name.lower()
    if "gbm" in folder or "idhwt" in folder:
        return {"grade_class": "high", "mutation": "IDH_wt"}
    if "oligo" in folder and "1p19q" in folder:
        return {"grade_class": "low" if "_g2" in folder else "high",
                "mutation": "IDH_mt_1p19q"}
    if "astro" in folder and "idhmt" in folder:
        return {"grade_class": "low" if "_g2" in folder else "high",
                "mutation": "IDH_mt"}
    return {"grade_class": "unknown", "mutation": "unknown"}


# ============================================================
# AGGREGATION
# ============================================================

def _mean_std(vals):
    if not vals:
        return ""
    m = float(np.mean(vals))
    s = float(np.std(vals, ddof=1)) if len(vals) > 1 else 0.0
    return f"{m:.4f} ± {s:.4f}"


def _write_testset_outputs(rows, run_dir):
    by_fold = defaultdict(list)
    for r in rows:
        by_fold[r["fold"]].append(r)

    # per-fold summary rows
    summary_header = ["fold", "n_slides", "grade_acc", "subtype_acc",
                      "mutant_vs_wildtype_acc", "end_to_end_acc"]
    summary = []
    fold_vals = defaultdict(list)
    for fold, rs in sorted(by_fold.items()):
        tumour = [r for r in rs if r.get("stage_b_evaluated") == 1]
        grade_acc = np.mean([r["grade_correct"] for r in rs]) if rs else 0
        sub_acc = np.mean([r["subtype_correct"] for r in tumour]) if tumour else 0
        mvw = np.mean([r["mutant_status_correct"] for r in tumour]) if tumour else 0
        e2e = np.mean([r["end_to_end_correct"] for r in rs]) if rs else 0
        summary.append([fold, len(rs), round(grade_acc, 4), round(sub_acc, 4),
                        round(mvw, 4), round(e2e, 4)])
        fold_vals["grade_acc"].append(grade_acc)
        fold_vals["subtype_acc"].append(sub_acc)
        fold_vals["mutant_vs_wildtype_acc"].append(mvw)
        fold_vals["end_to_end_acc"].append(e2e)
    summary.append(["mean ± std", "",
                    _mean_std(fold_vals["grade_acc"]),
                    _mean_std(fold_vals["subtype_acc"]),
                    _mean_std(fold_vals["mutant_vs_wildtype_acc"]),
                    _mean_std(fold_vals["end_to_end_acc"])])

    # 1p/19q codeletion table (true mutants only, per fold + overall)
    codel_header = ["fold", "n_true_mutant",
                    "codeleted_sensitivity", "noncodeleted_sensitivity",
                    "codeletion_acc_within_called_mutant"]
    codel = []
    cod_vals = defaultdict(list)
    for fold, rs in sorted(by_fold.items()):
        mut = [r for r in rs if r.get("true_mutation") in MUTANT_SET]
        if not mut:
            continue
        oligo = [r for r in mut if r["true_mutation"] == "IDH_mt_1p19q"]
        astro = [r for r in mut if r["true_mutation"] == "IDH_mt"]
        sens_cod = (np.mean([r["subtype_pred"] == "IDH_mt_1p19q" for r in oligo])
                    if oligo else float("nan"))
        sens_non = (np.mean([r["subtype_pred"] == "IDH_mt" for r in astro])
                    if astro else float("nan"))
        called = [r for r in mut if r.get("subtype_pred") in MUTANT_SET]
        acc_called = (np.mean([(r["subtype_pred"] == "IDH_mt_1p19q") ==
                               (r["true_mutation"] == "IDH_mt_1p19q") for r in called])
                      if called else float("nan"))
        codel.append([fold, len(mut), _r(sens_cod), _r(sens_non), _r(acc_called)])
        for k, v in (("codeleted_sensitivity", sens_cod),
                     ("noncodeleted_sensitivity", sens_non),
                     ("codeletion_acc_within_called_mutant", acc_called)):
            if not np.isnan(v):
                cod_vals[k].append(v)
    codel.append(["mean ± std", "",
                  _mean_std(cod_vals["codeleted_sensitivity"]),
                  _mean_std(cod_vals["noncodeleted_sensitivity"]),
                  _mean_std(cod_vals["codeletion_acc_within_called_mutant"])])

    extra = {"summary": (summary_header, summary),
             "codeletion_1p19q": (codel_header, codel)}
    write_xlsx_report(rows, run_dir / "e2e_testset_report.xlsx", extra_sheets=extra)


def _r(v):
    return "" if (isinstance(v, float) and np.isnan(v)) else round(float(v), 4)


# ============================================================
# CLI  (flags override CONFIG)
# ============================================================

def parse_args():
    c = CONFIG
    p = argparse.ArgumentParser()
    p.add_argument("--testset", action="store_true",
                   default=(c["run_mode"] == "testset"))
    p.add_argument("--run-name", default=c["run_name"])
    p.add_argument("--input", default=c["input"])
    p.add_argument("--model", default=c["model"])
    p.add_argument("--stage-a-ckpt", default=c["stage_a_ckpt"])
    p.add_argument("--stage-b-ckpt", default=c["stage_b_ckpt"])
    p.add_argument("--stage-a-glob", default=c["stage_a_glob"])
    p.add_argument("--stage-b-glob", default=c["stage_b_glob"])
    p.add_argument("--splits-dir", default=c["splits_dir"])
    p.add_argument("--sidecar-root", default=c["sidecar_root"])
    p.add_argument("--wsi-root", default=c["wsi_root"])
    p.add_argument("--control-image", default=c["control_image"])
    p.add_argument("--percentile", type=float, default=c["percentile"])
    p.add_argument("--cap", type=int, default=c["cap"])
    p.add_argument("--min-tissue", type=float, default=c["min_tissue"])
    p.add_argument("--target-mpp", type=float, default=c["target_mpp"])
    p.add_argument("--thumb-size", type=int, default=c["thumb_size"])
    p.add_argument("--output-size", type=int, default=c["output_size"])
    p.add_argument("--stage-b-riskmap", action="store_true",
                   default=c["stage_b_riskmap"])
    p.add_argument("--occlusion-topk", type=int, default=c["occlusion_topk"])
    p.add_argument("--device", default=c["device"])
    return p.parse_args()


def main():
    args = parse_args()
    if args.testset:
        run_testset(args)
    else:
        run_single_or_folder(args)


if __name__ == "__main__":
    main()